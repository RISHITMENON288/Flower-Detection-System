import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# --- Color Palette ---
COLOR_PRIMARY = RGBColor(0x1E, 0x3A, 0x8A)     # Deep Navy
COLOR_SECONDARY = RGBColor(0x02, 0x84, 0xC7)   # Slate Blue
COLOR_DARK = RGBColor(0x0F, 0x17, 0x2A)        # Charcoal Text
COLOR_MUTED = RGBColor(0x64, 0x74, 0x8B)       # Slate Gray
COLOR_ALERT = RGBColor(0xB4, 0x53, 0x09)       # Amber

HEX_HEADER_BG = "1E3A8A"                       # Dark Navy for Table Headers
HEX_ZEBRA_BG = "F8FAFC"                        # Light Slate for Alternating Rows
HEX_BORDER = "CBD5E1"                          # Slate Border
HEX_PRIMARY_BORDER = "1E3A8A"

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def format_table(table, col_widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Repeat header row & prevent row split across pages
    for i, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if i == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

    # Borders
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{HEX_PRIMARY_BORDER}"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>
            <w:insideV w:val="none"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

    for i, row in enumerate(table.rows):
        for j, cell in enumerate(row.cells):
            set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if col_widths and j < len(col_widths):
                cell.width = Inches(col_widths[j])

            # Header style
            if i == 0:
                set_cell_background(cell, HEX_HEADER_BG)
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for run in p.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        run.font.size = Pt(9.5)
                        run.font.name = 'Calibri'
            else:
                if i % 2 == 1:
                    set_cell_background(cell, HEX_ZEBRA_BG)
                else:
                    set_cell_background(cell, "FFFFFF")
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9.5)
                        run.font.name = 'Calibri'
                        run.font.color.rgb = COLOR_DARK

def add_callout(doc, text, title="NOTE / OBSERVATION"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{HEX_HEADER_BG}"/>
            <w:top w:val="none"/>
            <w:right w:val="none"/>
            <w:bottom w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"✦ {title}: ")
    r_title.bold = True
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(9.5)
    r_title.font.color.rgb = COLOR_PRIMARY

    r_text = p.add_run(text)
    r_text.italic = True
    r_text.font.name = 'Calibri'
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = COLOR_DARK
    doc.add_paragraph()

def add_placeholder_box(doc, fig_title, req_image_desc, evidence_source, fig_caption):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "FAFAFA")
    set_cell_margins(cell, top=300, bottom=300, left=250, right=250)

    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="dashed" w:sz="12" w:space="0" w:color="94A3B8"/>
            <w:top w:val="dashed" w:sz="12" w:space="0" w:color="94A3B8"/>
            <w:right w:val="dashed" w:sz="12" w:space="0" w:color="94A3B8"/>
            <w:bottom w:val="dashed" w:sz="12" w:space="0" w:color="94A3B8"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)

    r1 = p.add_run("[SCREENSHOT PLACEHOLDER — INSERT IMAGE HERE]\n")
    r1.bold = True
    r1.font.name = 'Calibri'
    r1.font.size = Pt(11)
    r1.font.color.rgb = COLOR_PRIMARY

    r2 = p.add_run(f"{fig_title}\n")
    r2.bold = True
    r2.font.name = 'Calibri'
    r2.font.size = Pt(10)
    r2.font.color.rgb = COLOR_DARK

    r3 = p.add_run(f"Required image: {req_image_desc}\n")
    r3.font.name = 'Calibri'
    r3.font.size = Pt(9)
    r3.font.color.rgb = COLOR_MUTED

    r4 = p.add_run(f"Evidence from: {evidence_source}\n(Click inside this box, delete this text, then Insert → Pictures. Keep caption below.)")
    r4.font.name = 'Calibri'
    r4.font.size = Pt(8.5)
    r4.font.italic = True
    r4.font.color.rgb = COLOR_MUTED

    # Caption
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(4)
    cp.paragraph_format.space_after = Pt(12)
    c_run = cp.add_run(fig_caption)
    c_run.font.name = 'Calibri'
    c_run.font.size = Pt(9.5)
    c_run.font.italic = True
    c_run.font.color.rgb = COLOR_DARK

def add_image_figure(doc, img_path, caption_text, width_inches=6.0):
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run()
        run.add_picture(img_path, width=Inches(width_inches))

        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_before = Pt(2)
        cp.paragraph_format.space_after = Pt(14)
        c_run = cp.add_run(caption_text)
        c_run.font.name = 'Calibri'
        c_run.font.size = Pt(9.5)
        c_run.font.italic = True
        c_run.font.color.rgb = COLOR_DARK
    else:
        add_placeholder_box(doc, caption_text, f"Image not found at {img_path}", "Project directory", caption_text)
